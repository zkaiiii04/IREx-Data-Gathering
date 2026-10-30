"""Generate the fixture workbooks for the Collection Workbook tests.

Every record here is invented. The sheet, table and column names copy the
real Main Database, so the macros meet the same structures they will meet
there. Never copy names or figures from the real workbooks into this file.

Run from anywhere:  python3 tests/fixtures/make_fixtures.py

Writes:
  MainDatabase.xlsx                         the Main Database "as it is today"
  check/IREX_Rehab_Data_Collection_CY2025.xlsx
                                            a Collection Workbook exported for
                                            2025, then edited by Rehab Staff
  check/IREX_Rehab_Data_Collection_CY2025 – ID lookup.xlsx
                                            the lookup written at that export
"""

import datetime as dt
import os

import openpyxl
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = os.path.dirname(os.path.abspath(__file__))
D = dt.datetime
NO_FORMULA = '=IF(INDEX(B:B,ROW()-1)="No.",1,INDEX(B:B,ROW()-1)+1)'
EN_DASH = "–"

# --------------------------------------------------------------------------
# Main Database headers (copied from the real workbook's header rows)
# --------------------------------------------------------------------------

PUB_HEADERS = [
    "No.", "Title of Article", "1st Author Name", "TTSH Contributor ", "Status",
    "Date of Publication\n\n", "pmid", "Journal Citation",
    "Associated Research/Innovation Project (If applicable)",
    "Full Author Listing (Optional)", "DOI", "Category of Article", "Other Remarks",
    "Journal Name", "Journal Category & JIF Percentile Listing \n(2025)",
    "JIF Percentile \n(2025)", "Journal Category & JIF Percentile Listing \n(2026)",
    "JIF Percentile \n(2026)",
]
PRES_HEADERS = [
    "No.", "Title of Presentation/Topic\n\n\n", "Full Name ", "Status",
    "Date of Presentation", "Start Date of Event", "End Date of Event",
    "Title of  Related Research, Innovation Projects /Programmes",
    "Name of Conference/ Symposium/ Seminar", "Name of Organisation",
    "Country where the conference/sympoisum/Seminar was held", "Other Remarks",
]
MEDIA_HEADERS = [
    "No.", "Full name", "Platform of Media\n\n", "Media Publishing Date \n\n",
    "Title of Media Published & Description \n\n(Briefly provide the description "
    "and or purpose of the published media)",
    "Name of Media Published Source \n\n(For eg The Straits Times or online eg "
    "ChannelnewsAsia.com) \n",
    "Website link to media/ Source\n\n(If applicable)\n",
    "Associated Project Title \n\n(If applicable)", "Other Remarks",
]
TD_HEADERS = [
    "Name", "Family Group", "Award Description", "Project Name", "LOA Start Date",
    "LOA End Date", "Grant Amount", "Year", "Category",
]
EDU_HEADERS = [
    "Date/Year/Period", "Year", "Type of Material/Event",
    "Area(s) of Focus (Professional vs Patients/Care Givers/General Public)",
    "Title of Course/Educational Materials/Educational Events/Outreach",
    "Family Group", "Trainer(s)/Speaker(s) ", "Collaboration Partner(s), if any",
    "Number of Participants, if applicable", "Local/Overseas", "Country (if overseas)",
    "Organisers",
]
PRJ_HEADERS = [
    "Grant Ref No", "Project Name", "Project Alias", "Grant Submitted Date",
    "Grant Approved Date", "Category", "Description", "Grant / Collaborator Partners",
    "Grant Group", "Grant Type", "Project Name & Grant/Partners", "PI",
    "Co-I/Collaborator", "PI/Co-I/Collaborator", "Admin/Ops Support", "RA Support",
    "LOA Start Date", "LOA End Date", "Tracking Timestamp (from 2021 onwards)",
    "LOA Start & End Date", "Grant Amount", "Grant Amt (Awarded to TTSH)",
    "Grant Amt & Awarded to TTSH", "Indirect Cost (Credited)", "Status",
    "Progress Update", "Progress Update (Archive)", "DSRB Ref No",
    "DSRB Submitted Date", "DSRB Approved Date",
    "Assigned Cost Centre (based On IO Code)", "Recruitment Target",
    "PA/RCA effective date ", "PA/RCA end date ", "IO Code Group", "IO Code",
    "Manpower", "Budget", "Utilisation", "Commitment", "Balance",
]
PRJ_STATUSES = [
    "Presented", "Published", "Pending", "Applied", "Under Review", "Awarded",
    "In Progress", "Completed", "Incomplete", "Dropped", "Applied (Oversight)",
    "Under Review (Oversight)", "Awarded (Oversight)", "In Progress (Oversight)",
    "Completed (Oversight)", "Incompleted (Oversight)", "Dropped (Oversight)",
]

# --------------------------------------------------------------------------
# Main Database records (invented). Comments say what each row tests for 2025.
# --------------------------------------------------------------------------

# No. comes from the formula; the comment gives the number it computes.
PUBS = [
    # No. 1: in 2025
    ["Robotic gait training after stroke: a pilot study", "Lim A", "Tan Wei Ling\nGoh Mei Xin",
     "Published", D(2025, 3, 10), 39000001, "J Test Rehabil. 2025;1:1-9", None, None,
     "10.1000/test.0001", "Original Research Paper", None, "Journal of Test Rehabilitation",
     "REHABILITATION (80%)", 80, "REHABILITATION (82%)", 82],
    # No. 2: 2024, not exported
    ["Falls prevention in the community", "Ho B", "Raj Kumar", "Published", D(2024, 11, 2),
     39000002, "Test Geriatr J. 2024;3:4-8", None, None, None, None, None,
     "Test Geriatrics Journal", "GERIATRICS (55%)", 55, "-", None],
    # No. 3: text date that the macro can read, in 2025
    ["Spasticity outcomes after botulinum toxin", "Tan Wei Ling", "Tan Wei Ling", "Published",
     "20 Aug 25", 39000003, "Test Neuro Rehab. 2025;2:10-19", None, None, None,
     "Original Research Paper", None, "Test Neuro Rehab", "NEUROSCIENCES (60%)", 60, "-", None],
    # No. 4: unreadable date "(blank)": exported and listed
    ["Caregiver burden in spinal cord injury", "Goh Mei Xin", "Goh Mei Xin", "Accepted",
     "(blank)", "-", "Test SCI J. In press", None, None, None, None, None, "Test SCI Journal",
     "-", None, "-", None],
    # No. 5: Excel date number stored as text (45700 = 12 Feb 2025)
    ["Upper limb telerehabilitation", "Raj Kumar", "Raj Kumar", "Published", "45700",
     39000005, "Test Telehealth. 2025;5:1-5", None, None, "10.1000/test.0005", None, None,
     "Test Telehealth", "HEALTH POLICY (40%)", 40, "-", None],
    # No. 6: in 2025, no DOI
    ["Dysphagia screening in acute stroke", "Lee Hui Min", "Lee Hui Min", "Under Review",
     D(2025, 6, 1), None, None, None, None, None, None, None, "Test Stroke Journal", "-", None,
     "-", None],
    # No. 7: 2026, not exported
    ["Cognitive rehabilitation trial protocol", "Ong C", "Tan Wei Ling", "Accepted",
     D(2026, 1, 5), None, "Test Trials. 2026", None, None, None, None, None, "Test Trials",
     "-", None, "-", None],
]

PRES = [
    # No. 1: in 2025
    ["Early mobilisation in intensive care", "Tan Wei Ling", "Presented", D(2025, 5, 20),
     D(2025, 5, 19), D(2025, 5, 21), None, "Test Rehab Congress 2025", None, "Singapore", None],
    # No. 2: in 2025, End Date "(blank)"
    ["Wheelchair skills training", "Goh Mei Xin", "Accepted", D(2025, 10, 3), D(2025, 10, 2),
     "(blank)", None, "Test Seating Symposium", None, None, None],
    # No. 3: 2023, not exported
    ["Pain coping strategies", "Raj Kumar", "Presented", D(2023, 8, 1), D(2023, 8, 1),
     D(2023, 8, 2), None, "Test Pain Meeting", None, None, None],
    # No. 4: in 2025; renamed since the export (the lookup holds the old title)
    ["Virtual reality balance training for older adults", "Raj Kumar", "Presented",
     D(2025, 2, 14), D(2025, 2, 13), D(2025, 2, 15), "VR Balance", "Test Falls Forum", None,
     "Singapore", None],
    # No No. (the cell is empty): 2022, not exported, but listed in the export report
    ["Return to work after stroke", "Lee Hui Min", "Presented", D(2022, 5, 5), D(2022, 5, 5),
     D(2022, 5, 6), None, "Test Work Rehab Day", None, None, None],
]
PRES_WITHOUT_NO = 5

MEDIA = [
    # No. 1-3: one story carried by three outlets, in 2025
    ["Stroke Rehab Programme", "Digital News (Online articles, etc)", D(2025, 4, 1),
     "New stroke rehab centre opens", "The Test Times", "https://example.com/a", None, None],
    ["Stroke Rehab Programme", "Newspaper", D(2025, 4, 1), "New stroke rehab centre opens",
     "The Test Times Print", None, None, None],
    ["Stroke Rehab Programme", "Radio", D(2025, 4, 2), "New stroke rehab centre opens",
     "Test Radio 99", "https://example.com/radio", None, None],
    # No. 4: 2024, not exported
    ["Tan Wei Ling", "Social Media", D(2024, 12, 1), "Therapy tips video", "TestTube",
     "https://example.com/v", None, None],
    # No. 5: unreadable date
    ["Goh Mei Xin", "Websites & Blogs", "Nov 26(KIV)", "Living well after spinal cord injury",
     "Test Blog", None, None, None],
]

TDS = [
    ["Tan Wei Ling", "AHS", "Test Fellowship Award", "Gait study", D(2025, 1, 1),
     D(2026, 12, 31), 50000, 2025, "Talent Development"],
    ["Goh Mei Xin", "Medical", "Test Teaching Award", "-", "-", "-", "-", 2025, "Special Award"],
    # 2024, not exported
    ["Raj Kumar", "AHS", "Test Research Award", "-", "-", "-", "-", 2024, "Special Award"],
    # Year empty: exported and listed
    ["Lee Hui Min", "Nursing", "Test Nursing Award", "-", "-", "-", "-", None, "Special Award"],
]

EDU_IN_TABLE = [
    [D(2025, 3, 15), 2025, "Workshop/Masterclass", "Professional",
     "Spasticity management workshop", None, "Tan Wei Ling", "Test Hospital", 40, "Local",
     "Singapore", "Rehab Dept"],
    # 2024, not exported
    [D(2024, 9, 10), 2024, "Webinar", "Patients/Care Givers/General Public",
     "Living with stroke webinar", None, "Goh Mei Xin", "Test Hospital", 120, "Local",
     "Singapore", None],
    # Year empty, date in 2025: exported
    [D(2025, 7, 1), None, "Webinar", "Professional", "Pressure injury prevention webinar",
     None, "Lee Hui Min", "Test Hospital", 60, "Local", "Singapore", None],
    # Year empty, date in 2024: not exported
    [D(2024, 7, 1), None, "Webinar", "Professional", "Seating clinic update", None,
     "Raj Kumar", "Test Hospital", 30, "Local", "Singapore", None],
    # Year 2025, no date: exported
    [None, 2025, "Presentation/Sharing", "Professional", "Caregiver training sharing", None,
     "Goh Mei Xin", "Test Hospital", None, "Local", "Singapore", None],
    # Year empty, unreadable date: exported and listed
    ["25 & 26 Feb 25", None, "Workshop/Masterclass", "Professional",
     "Hand therapy splinting course", None, "Tan Wei Ling", "Test Hospital", 25, "Local",
     "Singapore", None],
]
# Rows typed below tbl_Education after a gap, as in the real Main Database:
# a bold label row, then records.
EDU_BELOW_LABEL = "PARTNER PROGRAMME"
EDU_BELOW = [
    [D(2025, 8, 18), 2025, "Workshop/Masterclass", "Professional", "Community rehab workshop",
     None, "Raj Kumar", "Test Partner", 35, "Local", "Singapore", None],
    [D(2022, 8, 18), 2022, "Workshop/Masterclass", "Professional", "Old outreach talk", None,
     "Raj Kumar", "Test Partner", 20, "Local", "Singapore", None],
]


def project(ref, name, alias, category, partner, gtype, pi, coi, start, end, amount,
            ttsh, status):
    row = {h: None for h in PRJ_HEADERS}
    row.update({
        "Grant Ref No": ref, "Project Name": name, "Project Alias": alias,
        "Category": category, "Description": "Invented project for tests",
        "Grant / Collaborator Partners": partner, "Grant Group": "Test Group A",
        "Grant Type": gtype, "Project Name & Grant/Partners": f"{alias}\n({partner})",
        "PI": pi, "Co-I/Collaborator": coi, "PI/Co-I/Collaborator": f"{pi}\n{coi or ''}",
        "Admin/Ops Support": "Admin Person", "RA Support": "RA Person",
        "LOA Start Date": start, "LOA End Date": end, "Grant Amount": amount,
        "Grant Amt (Awarded to TTSH)": ttsh, "Indirect Cost (Credited)": "10%",
        "Status": status, "Progress Update": "On track", "Progress Update (Archive)": "-",
        "Assigned Cost Centre (based On IO Code)": "CC-TEST", "IO Code Group": "IOG-TEST",
        "IO Code": "IO-TEST-" + (ref or "NA"), "Manpower": "1 RA", "Budget": "100000",
        "Utilisation": "40000", "Balance": "60000",
    })
    return [row[h] for h in PRJ_HEADERS]


PRJS = [
    # spans 2025: exported
    project("GRT-001", "Robot-assisted gait study", "RoboGait", "Research", "Test Fund",
            "Extramural", "Tan Wei Ling", "Goh Mei Xin", D(2023, 4, 1), D(2026, 3, 31),
            500000, 300000, "In Progress"),
    # starts in 2025: exported
    project("GRT-002", "Home exercise app", "HomeEx", "Innovation", "Test Innovation Fund",
            "Intramural", "Goh Mei Xin", None, D(2025, 6, 1), D(2027, 5, 31), 80000, 80000,
            "In Progress"),
    # ended 2024: not exported
    project("GRT-003", "Splint materials comparison", "Splints", "Research", "Test Fund",
            "Intramural", "Raj Kumar", None, D(2021, 1, 1), D(2024, 12, 31), 20000, 20000,
            "Completed"),
    # starts 2026: not exported
    project("GRT-004", "Aquatic therapy trial", "Aqua", "Research", "Test Fund", "Extramural",
            "Lee Hui Min", None, D(2026, 2, 1), D(2028, 1, 31), 150000, 150000, "Awarded"),
    # no LOA End: exported
    project("GRT-005", "Community stroke registry", "Registry", "LF", "Test Foundation",
            "Extramural", "Tan Wei Ling", "Raj Kumar", D(2022, 1, 1), None, 250000, 250000,
            "In Progress"),
    # Applied, no LOA, no Grant Ref No: exported
    project(None, "Exoskeleton feasibility", "Exo", "Innovation", "Test Innovation Fund",
            "Intramural", "Raj Kumar", None, None, None, 60000, None, "Applied"),
    # Applied (Oversight), no LOA: exported
    project("GRT-007", "Regional falls network", "FallsNet", "Research", "Test Fund",
            "Extramural", "Goh Mei Xin", None, None, None, 90000, None,
            "Applied (Oversight)"),
    # In Progress, no LOA: not exported
    project("GRT-008", "Orphan project with no dates", "Orphan", "Research", "Test Fund",
            "Intramural", "Lee Hui Min", None, None, None, 10000, None, "In Progress"),
    # ends on 1 Jan 2025: exported
    project("GRT-009", "Pressure mapping pilot", "PMap", "Research", "Test Fund", "Intramural",
            "Lee Hui Min", None, D(2024, 3, 1), D(2025, 1, 1), 30000, 30000, "Completed"),
]


def add_table(ws, name, ref):
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)


def fixed_props(wb):
    stamp = dt.datetime(2026, 1, 1)
    wb.properties.created = stamp
    wb.properties.modified = stamp
    wb.properties.creator = "IREx test fixtures"


def make_main_database(path):
    wb = openpyxl.Workbook()
    lst = wb.active
    lst.title = "tbl_List"
    lst["H1"] = "Project Status"
    for i, s in enumerate(PRJ_STATUSES, start=2):
        lst.cell(i, 8, s)
    add_table(lst, "tbllist_PrjStatus", f"H1:H{len(PRJ_STATUSES) + 1}")
    wb.defined_names["list_PrjStatus"] = DefinedName(
        "list_PrjStatus", attr_text="tbllist_PrjStatus[Project Status]")

    wb.create_sheet("Dashboard")["B2"] = "Team dashboard (team only)"

    prj = wb.create_sheet("DB_PrjGeneralInfo")
    prj.append(PRJ_HEADERS)
    for r in PRJS:
        prj.append(r)
    add_table(prj, "tbl_PrjGeneralInfo", f"A1:AO{len(PRJS) + 1}")
    status_col = PRJ_HEADERS.index("Status") + 1
    dv = DataValidation(type="list", formula1="list_PrjStatus", allow_blank=True)
    dv.add(f"Y2:Y{len(PRJS) + 1}")
    prj.add_data_validation(dv)
    assert status_col == 25
    for r in range(2, len(PRJS) + 2):
        for c in (17, 18):
            prj.cell(r, c).number_format = "dd\\ mmm\\ yy"
        for c in (21, 22):
            prj.cell(r, c).number_format = '"$"#,##0'

    wb.create_sheet("Gantt Chart")["A1"] = "Gantt chart (team only)"
    g = wb.create_sheet("DB_PrjGanttChart")
    g.sheet_state = "hidden"
    g.append(["Task", "Start"])
    g.append(["Plan", D(2025, 1, 1)])
    wb.create_sheet("DB_ManpowerRate").append(["Grade", "Monthly Rate"])
    wb.create_sheet("DB_RA_Distrib").append(["RA", "Project"])

    edu = wb.create_sheet("DB_Education")
    for c, h in enumerate(EDU_HEADERS, start=2):
        edu.cell(1, c, h)
    for i, rec in enumerate(EDU_IN_TABLE, start=2):
        for c, v in enumerate(rec, start=2):
            edu.cell(i, c, v)
    last = len(EDU_IN_TABLE) + 1
    add_table(edu, "tbl_Education", f"B1:M{last}")
    label_row = last + 6          # a gap of blank rows, as in the real sheet
    edu.cell(label_row, 2, EDU_BELOW_LABEL).font = openpyxl.styles.Font(bold=True)
    for i, rec in enumerate(EDU_BELOW, start=label_row + 1):
        for c, v in enumerate(rec, start=2):
            edu.cell(i, c, v)
    end = label_row + len(EDU_BELOW)
    for col, f1 in (("D", '"Workshop/Masterclass, Webinar, Presentation/Sharing"'),
                    ("E", '"Professional, Patients/Care Givers/General Public"'),
                    ("K", '"Local, Overseas"')):
        dv = DataValidation(type="list", formula1=f1, allow_blank=True)
        dv.add(f"{col}2:{col}{end}")
        edu.add_data_validation(dv)

    imp = wb.create_sheet("Impact Factor-2026")
    imp.append(["Journal Name", "Abbreviated Journal", "Percentile"])
    imp.append(["Journal of Test Rehabilitation", "J Test Rehabil", 82])
    add_table(imp, "tbl_ImpactFactor_2026", "A1:C2")

    pres = wb.create_sheet("DB_Poster& Oral Presentation")
    for c, h in enumerate(PRES_HEADERS, start=2):
        pres.cell(1, c, h)
    for i, rec in enumerate(PRES, start=2):
        if i - 1 != PRES_WITHOUT_NO:
            pres.cell(i, 2, NO_FORMULA)
        for c, v in enumerate(rec, start=3):
            pres.cell(i, c, v)
    add_table(pres, "tbl_PosterOralPresentation", f"B1:M{len(PRES) + 1}")
    pres.column_dimensions["I"].hidden = True

    pub = wb.create_sheet("DB_Publications")
    pub["O1"], pub["P1"], pub["Q1"] = "Input Start Year", 2025, '="Top 25% Count: "&COUNTIF(Q4:Q20,">=75")'
    pub["O2"], pub["P2"], pub["Q2"] = "Input End Year", 2025, '="Top 10% Count: "&COUNTIF(Q4:Q20,">=90")'
    for c, h in enumerate(PUB_HEADERS, start=2):
        pub.cell(3, c, h)
    for i, rec in enumerate(PUBS, start=4):
        pub.cell(i, 2, NO_FORMULA)
        for c, v in enumerate(rec, start=3):
            pub.cell(i, c, v)
        pub.cell(i, 7).number_format = "d-mmm-yy"
    add_table(pub, "tbl_Publication", f"B3:S{len(PUBS) + 3}")
    for col in ("J", "K", "M", "N"):
        pub.column_dimensions[col].hidden = True
    pub.column_dimensions["C"].width = 51

    media = wb.create_sheet("DB_Media")
    for c, h in enumerate(MEDIA_HEADERS, start=2):
        media.cell(1, c, h)
    for i, rec in enumerate(MEDIA, start=2):
        media.cell(i, 2, NO_FORMULA)
        for c, v in enumerate(rec, start=3):
            media.cell(i, c, v)

    td = wb.create_sheet("DB_Talent Develop")
    for c, h in enumerate(TD_HEADERS, start=2):
        td.cell(1, c, h)
    for i, rec in enumerate(TDS, start=2):
        for c, v in enumerate(rec, start=2):
            td.cell(i, c, v)
    add_table(td, "tbl_TalentDevelop", f"B1:J{len(TDS) + 1}")
    # A pivot-style summary to the right of the table, as in the real sheet.
    td["L2"], td["M2"], td["N2"] = "Row Labels", "Name", "Award Description"
    td["L3"] = "Special Award"
    td["L4"], td["M4"], td["N4"] = 2025, "Goh Mei Xin", "Test Teaching Award"

    rie = wb.create_sheet("DB_RIE Namelist")
    rie.append([None, "Name", "Family Group"])
    rie.append([None, "Tan Wei Ling", "Rehab AHS"])
    wb.create_sheet("Visualisation-AA")["B2"] = "Charts (team only)"

    fixed_props(wb)
    wb.save(path)


# --------------------------------------------------------------------------
# Collection Workbook fixture for the check: exported for 2025, then edited.
# Layout follows Export Staff Copy: note in A1, headers on row 3 from column A,
# records from row 4, hidden Export ID as the last exported column.
# --------------------------------------------------------------------------

NOTE = "How to find your records: press Ctrl+F and search for your name, or use the filter arrow."
HEADER_ROW = 3


def pub_cw(no, rec):
    # Every Publications column except the four JIF columns, with No. as a value.
    return [no] + rec[:13]


def write_cw_tab(wb, title, headers, rows, table_rows, extra=None):
    """rows: list of lists (None for a blank row). table_rows: how many of the
    rows sit inside the Excel table. extra: {(row_offset, header): value}."""
    ws = wb.create_sheet(title)
    ws["A1"] = NOTE
    all_headers = headers + ["Export ID"]
    for c, h in enumerate(all_headers, start=1):
        ws.cell(HEADER_ROW, c, h)
    for i, rec in enumerate(rows):
        if rec is None:
            continue
        for c, v in enumerate(rec, start=1):
            ws.cell(HEADER_ROW + 1 + i, c, v)
    last_col = openpyxl.utils.get_column_letter(len(all_headers))
    add_table(ws, "Collection_" + title.split()[0].replace("&", ""),
              f"A{HEADER_ROW}:{last_col}{HEADER_ROW + table_rows}")
    ws.column_dimensions[last_col].hidden = True
    for (r, c), v in (extra or {}).items():
        ws.cell(r, c, v)
    return ws


def make_collection_workbook(folder):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    pub_headers = [PUB_HEADERS[0]] + PUB_HEADERS[1:14]
    p = {i + 1: rec for i, rec in enumerate(PUBS)}
    renamed_spaces = list(p[3])
    renamed_spaces[0] = "SPASTICITY  outcomes after botulinum toxin "   # case and spaces only
    renamed_spaces[2] = "Tan Wei Ling\r\n"                               # line break only
    renamed_spaces[4] = D(2025, 8, 20)                                    # real date vs "20 Aug 25"
    status_change = list(p[6])
    status_change[3] = "Published"
    blank_date = list(p[4])
    blank_date[4] = None                                                  # empty vs "(blank)"
    typed_update = list(p[6])
    typed_update[3] = "Accepted"
    pub_rows = [
        pub_cw(6, status_change) + ["PUB-0005"],                 # changed status
        pub_cw(1, p[1]) + ["PUB-0001"],                          # unchanged
        pub_cw(3, renamed_spaces) + ["PUB-0002"],                # formatting-only differences
        pub_cw(4, blank_date) + ["PUB-0003"],                    # empty vs (blank)
        # PUB-0004 (No. 5) was deleted by a member of Rehab Staff
        [None, "Hand therapy outcomes registry", "Ng Q", "Ng Qi Hui", "Accepted",
         D(2025, 9, 1), None, None, None, None, None, None, None, "Test Hand Journal", None],
        [None] * 3 + ["   "] + [None] * 11,                      # spaces only: ignored
        [None, "Orthotics survey", "Ng Q", "Ng Qi Hui", "Published", D(2025, 2, 2), None,
         None, None, None, None, None, None, "Test Orthotics", "PUB-9999"],   # unknown Export ID
        None,
        None,
        [None, "Prosthetic training survey", "Ng Q", "Ng Qi Hui", "Under Review",
         D(2025, 10, 10), None, None, None, None, None, None, None, "Test Prosthetics", None],
        pub_cw(None, typed_update) + [None],                     # status update typed as a new row
    ]
    ws = write_cw_tab(wb, "Publications", pub_headers, pub_rows, table_rows=7)
    ws.cell(HEADER_ROW, 17, "My notes")                          # extra column added by staff
    ws.cell(HEADER_ROW + 2, 17, "check with co-author")
    ws.cell(HEADER_ROW + 3, 20, "see email")                     # value outside the columns

    pr = {i + 1: rec for i, rec in enumerate(PRES)}
    old_title = list(pr[4])
    old_title[0] = "Virtual reality balance training"
    blank_end = list(pr[2])
    blank_end[5] = None
    write_cw_tab(wb, "Poster & Oral Presentation", [h for h in PRES_HEADERS], [
        [1] + pr[1] + ["PRES-0001"],
        [2] + blank_end + ["PRES-0002"],
        [4] + old_title + ["PRES-0003"],
    ], table_rows=3)

    m = {i + 1: rec for i, rec in enumerate(MEDIA)}
    link_change = list(m[3])
    link_change[5] = "https://example.com/radio-replay"
    new_outlet = list(m[1])
    new_outlet[4] = "Test News Online"
    retyped = list(m[1])
    write_cw_tab(wb, "Media", MEDIA_HEADERS, [
        [1] + m[1] + ["MED-0001"],
        [2] + m[2] + ["MED-0002"],
        [3] + link_change + ["MED-0003"],
        [5] + m[5] + ["MED-0004"],
        [None] + new_outlet + [None],                            # same story, new outlet: new
        [None] + retyped + [None],                               # same story and outlet: already recorded
    ], table_rows=4)

    t = {i + 1: rec for i, rec in enumerate(TDS)}
    amount_text = list(t[1])
    amount_text[6] = "50,000"
    year_added = list(t[4])
    year_added[7] = 2025
    write_cw_tab(wb, "Talent Development", TD_HEADERS, [
        amount_text + ["TD-0001"],
        t[2] + ["TD-0002"],
        year_added + ["TD-0003"],
    ], table_rows=3)

    e = EDU_IN_TABLE
    old_edu = list(e[0])
    old_edu[4] = "Spasticity workshop"
    write_cw_tab(wb, "Education", EDU_HEADERS, [
        old_edu + ["EDU-0001"],                                  # renamed in the Main Database since
        e[2] + ["EDU-0002"],
        e[4] + ["EDU-0003"],
        e[5] + ["EDU-0004"],
        [EDU_BELOW_LABEL] + [None] * 11 + ["EDU-0005"],
        EDU_BELOW[0] + ["EDU-0006"],
        [D(2025, 11, 5), 2025, "Webinar", "Professional", "Falls prevention webinar", None,
         "Ng Qi Hui", None, 80, "Local", "Singapore", None, None],          # new
    ], table_rows=7)

    prj_cols = ["Grant Ref No", "Project Name", "Category", "Grant / Collaborator Partners",
                "Grant Type", "PI", "Co-I/Collaborator", "LOA Start Date", "LOA End Date",
                "Grant Amount", "Grant Amt (Awarded to TTSH)", "Status"]
    idx = [PRJ_HEADERS.index(h) for h in prj_cols]
    exported = [PRJS[0], PRJS[1], PRJS[4], PRJS[5], PRJS[6], PRJS[8]]
    rows = []
    for n, rec in enumerate(exported, start=1):
        row = [rec[i] for i in idx] + [f"PRJ-{n:04d}"]
        rows.append(row)
    rows[0][prj_cols.index("Status")] = "Completed"              # changed status
    write_cw_tab(wb, "Projects", prj_cols, rows, table_rows=len(rows))

    fixed_props(wb)
    wb.save(os.path.join(folder, "IREX_Rehab_Data_Collection_CY2025.xlsx"))

    lk = openpyxl.Workbook()
    ws = lk.active
    ws.title = "ID lookup"
    ws.append(["Export ID", "Tab", "No.", "Title", "Key 1", "Key 2"])
    for n, (no, rec) in enumerate([(1, p[1]), (3, p[3]), (4, p[4]), (5, p[5]), (6, p[6])], 1):
        ws.append([f"PUB-{n:04d}", "Publications", no, rec[0], rec[9], rec[0]])
    for n, (no, title) in enumerate([(1, pr[1][0]), (2, pr[2][0]),
                                     (4, "Virtual reality balance training")], 1):
        ws.append([f"PRES-{n:04d}", "Poster & Oral Presentation", no, title, title, None])
    for n, (no, rec) in enumerate([(1, m[1]), (2, m[2]), (3, m[3]), (5, m[5])], 1):
        ws.append([f"MED-{n:04d}", "Media", no, rec[3], rec[3], rec[4]])
    for n, rec in enumerate([t[1], t[2], t[4]], 1):
        ws.append([f"TD-{n:04d}", "Talent Development", None, rec[2], rec[0], rec[2]])
    edu_exported = [old_edu, e[2], e[4], e[5], [EDU_BELOW_LABEL] + [None] * 11, EDU_BELOW[0]]
    for n, rec in enumerate(edu_exported, 1):
        ws.append([f"EDU-{n:04d}", "Education", None, rec[4], rec[4], rec[1]])
    for n, rec in enumerate(exported, 1):
        ws.append([f"PRJ-{n:04d}", "Projects", None, rec[1], rec[0], rec[1]])
    fixed_props(lk)
    lk.save(os.path.join(folder, f"IREX_Rehab_Data_Collection_CY2025 {EN_DASH} ID lookup.xlsx"))


if __name__ == "__main__":
    make_main_database(os.path.join(HERE, "MainDatabase.xlsx"))
    check = os.path.join(HERE, "check")
    os.makedirs(check, exist_ok=True)
    make_collection_workbook(check)
    print("fixtures written to", HERE)
